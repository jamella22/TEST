#!/usr/bin/env python3
import sys
import os
from typing import List

# DOCX
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# PDF
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib import colors


def is_heading(line: str) -> bool:
    text = line.strip()
    return text.isupper() and 2 <= len(text) <= 40


def read_lines(path: str) -> List[str]:
    with open(path, 'r', encoding='utf-8', errors='ignore') as f:
        return [ln.rstrip('\n') for ln in f.readlines()]


def export_docx(lines: List[str], out_path: str) -> None:
    doc = Document()
    # Margins
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # Base font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    # Ensure eastAsia font is set for compatibility
    try:
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    except Exception:
        pass

    first_non_empty_used = False
    for idx, raw in enumerate(lines):
        line = raw.rstrip()
        if line.strip() == '':
            doc.add_paragraph('')
            continue

        if not first_non_empty_used:
            p = doc.add_paragraph()
            run = p.add_run(line.strip())
            run.bold = True
            run.font.size = Pt(18)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            first_non_empty_used = True
            continue

        if is_heading(line):
            p = doc.add_paragraph()
            run = p.add_run(line.strip())
            run.bold = True
            run.font.size = Pt(13)
            continue

        if line.lstrip().startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(line.lstrip()[2:])
        else:
            doc.add_paragraph(line)

    doc.save(out_path)


def export_pdf(lines: List[str], out_path: str) -> None:
    doc = SimpleDocTemplate(
        out_path,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    styles = getSampleStyleSheet()

    # Customize styles
    styles.add(ParagraphStyle(
        name='CVTitle',
        parent=styles['Title'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        alignment=TA_LEFT,
        spaceAfter=8,
    ))
    styles.add(ParagraphStyle(
        name='CVHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12.5,
        leading=16,
        textColor=colors.black,
        spaceBefore=8,
        spaceAfter=4,
    ))
    styles.add(ParagraphStyle(
        name='CVBody',
        parent=styles['BodyText'],
        fontName='Helvetica',
        fontSize=10.5,
        leading=14,
        spaceAfter=2,
    ))
    styles.add(ParagraphStyle(
        name='CVBullet',
        parent=styles['BodyText'],
        fontName='Helvetica',
        fontSize=10.5,
        leading=14,
        leftIndent=12,
        bulletIndent=0,
        spaceAfter=1,
    ))

    story = []
    first_non_empty_used = False

    for idx, raw in enumerate(lines):
        line = raw.rstrip()
        if line.strip() == '':
            story.append(Spacer(1, 6))
            continue

        if not first_non_empty_used:
            story.append(Paragraph(line.strip(), styles['CVTitle']))
            first_non_empty_used = True
            continue

        if is_heading(line):
            story.append(Paragraph(line.strip(), styles['CVHeading']))
            continue

        if line.lstrip().startswith('- '):
            story.append(Paragraph(line.lstrip()[2:], styles['CVBullet'], bulletText='•'))
        else:
            story.append(Paragraph(line, styles['CVBody']))

    doc.build(story)


def process_file(path: str) -> None:
    base, ext = os.path.splitext(path)
    docx_out = base + '.docx'
    pdf_out = base + '.pdf'

    lines = read_lines(path)

    export_docx(lines, docx_out)
    export_pdf(lines, pdf_out)

    print(f"Exported: {docx_out} and {pdf_out}")


def main(argv: List[str]) -> int:
    if len(argv) < 2:
        print('Usage: export_cv.py <input1.txt> [<input2.txt> ...]')
        return 1
    for in_path in argv[1:]:
        if not os.path.isfile(in_path):
            print(f"Warning: {in_path} not found, skipping")
            continue
        try:
            process_file(in_path)
        except Exception as e:
            print(f"Error processing {in_path}: {e}")
            return 2
    return 0


if __name__ == '__main__':
    raise SystemExit(main(sys.argv))
